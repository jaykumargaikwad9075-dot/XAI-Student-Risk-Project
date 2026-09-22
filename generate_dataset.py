"""Generate a 200-student synthetic dataset that follows the same structure and
Risk Score Rules as the supplied 250-row template workbook."""
import numpy as np, pandas as pd, openpyxl, copy, sys
from openpyxl.styles import Font
SRC = "/mnt/user-data/uploads/Explainable_AI_Personalised_Student_Academic_Risk_Prediction_Dataset_Final__1_.xlsx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "data/Student_Dataset_200.xlsx"
rng = np.random.default_rng(2026)
N = 200

female = ["Aarohi","Aditi","Ankita","Anushka","Apeksha","Bhakti","Chaitali","Dipali","Gayatri","Harshada","Janhavi","Komal","Madhuri","Manasi","Mrunal","Nikita","Payal","Prachi","Rutuja","Sayali","Shreya","Snehal","Sonali","Swara","Tejaswini","Vidya","Yashashri","Pratiksha","Ritika","Shruti"]
male = ["Abhishek","Aditya","Akash","Ajinkya","Chinmay","Darshan","Harsh","Kedar","Mayur","Nikhil","Omkar","Pranav","Prathamesh","Rahul","Rohan","Sagar","Sahil","Shubham","Siddharth","Swapnil","Tejas","Utkarsh","Vaibhav","Vedant","Vikram","Yash","Sumit","Tushar","Onkar","Nilesh"]
last = ["Patil","Jadhav","Shinde","Pawar","Deshmukh","Chavan","More","Bhosale","Kadam","Kulkarni","Joshi","Thorat","Salunkhe","Mane","Sawant","Wagh","Dhumal","Lokhande","Suryawanshi","Kale","Gore","Yadav","Sharma","Gupta","Naik","Kumar","Singh","Verma","Gaikwad","Deshpande","Bhide","Kamble","Nimbalkar","Ghule","Phadtare"]

gender = rng.permutation(["Female"]*90 + ["Male"]*110)
names, used = [], set()
for g in gender:
    while True:
        n = f"{rng.choice(female if g=='Female' else male)} {rng.choice(last)}"
        if n not in used: used.add(n); break
    names.append(n)

# latent ability drives everything so features are realistically correlated
z = rng.normal(0, 1, N) - 0.15
def clip(x, lo, hi): return int(np.clip(round(x), lo, hi))
subj = {s: [clip(55 + 12*zi + rng.normal(0, 7), 20, 95) for zi in z]
        for s in ["Mathematics","Programming","Database","Statistics","AI"]}
att   = [clip(77 + 9*zi + rng.normal(0, 6), 50, 98) for zi in z]
study = [clip(3.6 + 1.3*zi + rng.normal(0, 0.9), 1, 6) for zi in z]
def behaviour(zi):
    s = zi + rng.normal(0, 0.55)
    return "Poor" if s < -0.35 else "Average" if s < 0.25 else "Good" if s < 1.0 else "Excellent"
beh = [behaviour(zi) for zi in z]
assign = ["Complete" if (zi + rng.normal(0, 0.6)) > -0.3 else "Incomplete" for zi in z]
other  = ["Present" if (zi*0.5 + rng.normal(0, 1)) > -0.7 else "Not Present" for zi in z]

df = pd.DataFrame({
    "Student ID": [f"STU{i:03d}" for i in range(1, N+1)],
    "Student Name": names,
    "Student Email": [f"{n.lower().replace(' ','.')}{i}@studentmail.com" for i, n in enumerate(names, 1)],
    "Student Phone Number": [int(f"{rng.integers(7,10)}{rng.integers(0,10**9):09d}") for _ in range(N)],
    "Student Gender": gender,
    "Mathematics Marks": subj["Mathematics"], "Programming Marks": subj["Programming"],
    "Database Marks": subj["Database"], "Statistics Marks": subj["Statistics"], "AI Marks": subj["AI"],
    "Student Behaviour": beh, "Assignment Record": assign, "Other Activities": other,
    "Attendance (%)": att, "Study Hours (per day)": study,
})

# Risk Score Rules (template sheet 'Risk Score Rules'; inferred bands 0-2 Low, 3-5 Medium, 6+ High)
def band(x, b): return 3 if x < b[0] else 2 if x < b[1] else 1 if x < b[2] else 0
marks = ["Mathematics Marks","Programming Marks","Database Marks","Statistics Marks","AI Marks"]
score = (df[marks].mean(axis=1).apply(lambda x: band(x, [40,55,65]))
         + df["Attendance (%)"].apply(lambda x: band(x, [65,75,85]))
         + df["Study Hours (per day)"].apply(lambda h: 2 if h <= 1 else 1 if h == 2 else 0)
         + df["Student Behaviour"].map({"Poor":2,"Average":1,"Good":0,"Excellent":0})
         + (df["Assignment Record"] == "Incomplete")*2 + (df["Other Activities"] == "Not Present")*1)
df["Academic Risk"] = np.where(score <= 2, "Low Risk", np.where(score <= 5, "Medium Risk", "High Risk"))
df.to_csv("data/student_dataset_200.csv", index=False)
print(df["Academic Risk"].value_counts(), df["Student Gender"].value_counts(), sep="\n")
print(pd.crosstab(df["Student Behaviour"], df["Academic Risk"]))

wb = openpyxl.load_workbook(SRC)
ws = wb["Student Dataset"]
styles = {}
for c in range(1, 17):
    t = ws.cell(2, c)
    styles[c] = (copy.copy(t.font), copy.copy(t.fill), copy.copy(t.alignment), copy.copy(t.border), t.number_format)
ws.delete_rows(2, ws.max_row)
for r, row in enumerate(df.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(r, c, v.item() if hasattr(v, "item") else v)
        cell.font, cell.fill, cell.alignment, cell.border, cell.number_format = styles[c]
ws.auto_filter.ref = f"A1:P{N+1}"

sm = wb.create_sheet("Summary")
hf, hfill = copy.copy(ws["A1"].font), copy.copy(ws["A1"].fill)
f = Font(name="Arial", size=10)
for c, t in zip("ABC", ["Academic Risk", "Students", "Share"]):
    sm[f"{c}1"] = t; sm[f"{c}1"].font = copy.copy(hf); sm[f"{c}1"].fill = copy.copy(hfill)
for i, lab in enumerate(["Low Risk","Medium Risk","High Risk"], start=2):
    sm.cell(i,1,lab).font = f
    sm.cell(i,2,f"=COUNTIF('Student Dataset'!$P$2:$P${N+1},A{i})").font = f
    sm.cell(i,3,f"=B{i}/$B$5").font = f; sm.cell(i,3).number_format = "0.0%"
sm["A5"], sm["B5"], sm["C5"] = "Total", "=SUM(B2:B4)", "=SUM(C2:C4)"
for c in ("A5","B5","C5"): sm[c].font = Font(name="Arial", size=10, bold=True)
sm["C5"].number_format = "0.0%"
labels = ["Average marks (all 5 subjects)", "Average attendance (%)", "Average study hours / day", "Assignments incomplete (%)"]
forms = [f"=AVERAGE('Student Dataset'!F2:J{N+1})", f"=AVERAGE('Student Dataset'!N2:N{N+1})",
         f"=AVERAGE('Student Dataset'!O2:O{N+1})", f"=COUNTIF('Student Dataset'!L2:L{N+1},\"Incomplete\")/B5"]
for r, (l, fm) in enumerate(zip(labels, forms), start=7):
    sm.cell(r,1,l).font = f; sm.cell(r,2,fm).font = f; sm.cell(r,2).number_format = "0.0"
sm["B10"].number_format = "0.0%"
sm["A12"] = "Note: synthetic data generated with the template's Risk Score Rules (0-2 Low, 3-5 Medium, 6+ High). Not real student records."
sm["A12"].font = Font(name="Arial", size=9, italic=True)
sm.column_dimensions["A"].width = 34; sm.column_dimensions["B"].width = 12; sm.column_dimensions["C"].width = 10
wb.save(OUT)
